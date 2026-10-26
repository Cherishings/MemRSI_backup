from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
OUT=Path('outputs/paper-experiments');OUT.mkdir(exist_ok=True)
data=json.loads(Path('work/experiment-tables/v4-results.json').read_text())
pages=[]
def page(title,intro,caption,headers,rows,widths,notes):
 pages.append(dict(title=title,intro=intro,caption=caption,headers=headers,rows=rows,widths=widths,notes=notes))
page('实验安排与论文结构',
 '实验围绕三个问题展开：Guidance 和 Update 是否改善跨会话任务执行，自动更新规则是否进一步带来收益，以及冻结后的方法能否迁移到其他环境。下列表格保留已登记的比较；待测单元格表示尚无正式测试结果，不代表零分。',
 '表 1 实验问题与证据安排',
 ['实验','核心比较','报告内容','当前阶段'],[
 ['模块消融','完整方法、仅 G、仅 U、原生 PAOS','完整任务成功率；配对差值','开发轮已完成；正式测试待运行'],
 ['额外思考对照','完整方法 vs 普通思考加 U','相同调用机会下的收益；实际成本','开发轮已完成；正式测试待运行'],
 ['四类行为诊断','中断恢复、改序、取消、最终目标','各类行为正确率；首次违规记录','任务中已设置；分类汇总待完成'],
 ['自动 RSI','父规则、仅更新 G、仅更新 U、冻结选中规则','训练增益、验证门槛、独立物理结果','候选通过审查；回放阶段'],
 ['跨环境比较','LIBERO、Casa、Twin、两个额外 benchmark','各环境自身的任务指标','LIBERO 优先；其余待部署和登记'],
 ['效率与可靠性','所有不同规则组','调用、tokens、耗时、等待、失败分类','已有开发日志；正式表待测']
 ],[1.05,2.05,1.85,1.95],[
 '论文组织建议：Experimental Setup 说明表 2 和表 3；Main Results 使用表 4；Behavioral Diagnostics 使用表 5；RSI Analysis 使用表 6；Transfer Evaluation 使用表 7；Efficiency and Reliability 使用表 8。表 9 和表 10 是已完成的开发结果，应独立放在开发分析或附录。',
 '截至 2026 年 10 月 24 日，正式物理测试尚未启动。当前开发数据不能替代测试集，也不能据此宣布 G、U 或 RSI 的最终优势。'])
page('实验设置与数据划分',
 '我们在选定的 LIBERO 官方物理任务之上构建长程、多会话任务。任务包含桌面物品放置、厨房工作插入、顺序更新和任务取消。完整执行路径包含 20 个会话及物理执行段、19 次会话切换和 7 个计分检查点。',
 '表 2 当前设置与已登记的测试规模',
 ['项目','设置','说明'],[
 ['研究范围','基于 LIBERO 的长程多会话任务','不等同于 LIBERO 全套官方基准'],
 ['物理任务','选定 6 个官方任务','包括摆放杯子与布丁、放置碗、操作抽屉和微波炉等'],
 ['训练来源','4 种任务历史 × 每种 2 条件 = 8 条','父规则轨迹；可重建且选入回放的为 4 条'],
 ['保留验证来源','2 种任务历史 × 每种 2 条件 = 4 条','3 条可重建，固定选择 2 条；不供优化器提案'],
 ['正式物理测试','12 种任务历史 × 2 初始状态 × 2 厨房变体 = 48 条件','每个不同的最终规则组使用相同条件'],
 ['最多正式运行数','8 组 × 48 = 384 个完整任务','相同规则哈希合并为别名，实际可少于 384'],
 ['语言模型','gpt-5.6-luna，经本机网关','所有实验节点通过本机调用'],
 ['机器人策略','固定 pi05_libero CUDA/JAX 服务','同一配对的各组共享后端与 paired seed'],
 ['持久记忆上限','1024 tokens','MEMORY.md 与 HISTORY.md 合计'],
 ['活动输入上限','12288 tokens','包含消息、工具 schema、记忆及指导等'],
 ['每次输出上限','8192 tokens','按实际请求与新登记值报告'],
 ['每个任务调用上限','480 次','额外思考组每次决策增加一个指导形式调用'],
 ['学习器预算','24000 次逻辑调用总上限','延续上一版已用 2 次；提案输出仍为 8192'],
 ['学习器输入上限','49152 tokens','与执行 actor 的 12288 上限分开报告']
 ],[1.35,2.75,2.8],[
 '训练、验证、测试按任务历史划分。正式测试还使用登记的初始状态 20 和 21；训练与验证使用 0 和 1。物理任务和部分物理状态存在历史接触，因此只能表述为“与本轮学习隔离的任务组合”，不能笼统称为从未见过的环境。',
 '预算勘误：旧 budget.json 中 output_tokens_per_call 仍为 2048，但新 output_tokens 与实际请求 max_tokens 均为 8192。旧来源登记的 learning_call_limit 为 12000，当前学习控制器冻结值为 24000；论文应明确版本，不能混用。'])
page('方法组与消融含义',
 '令 G₀、U₀ 表示学习前的人工父规则，Gf、Uf 表示经过训练及保留验证选择后冻结的规则。若某算子没有被接受的更新，其冻结规则就等于父规则。关闭自定义 U 时，PAOS 仍保留原生归档机制。',
 '表 3 已登记的八个正式比较组',
 ['论文名称','代码组名','Guidance','Update','回答的问题'],[
 ['原生 PAOS','native','关闭','原生归档','基础能力'],
 ['仅 G','g_only','Gf','原生归档','仅指导是否有效'],
 ['仅 U','u_only','关闭','Uf','仅更新记忆是否有效'],
 ['学习前 G＋U','gu_parent','G₀','U₀','固定父规则基线'],
 ['冻结后 G＋U','gu_learned','Gf','Uf','选中规则的完整表现'],
 ['仅更新 G 的 RSI','rsi_g_only','Gf','U₀','G 规则变化的贡献'],
 ['仅更新 U 的 RSI','rsi_u_only','G₀','Uf','U 规则变化的贡献'],
 ['普通思考＋U','think_u','普通额外思考','Uf','收益是否仅来自额外模型调用']
 ],[1.55,1.25,.8,.9,2.4],[
 '“仅 G／仅 U”关闭另一个自定义模块；“仅更新 G／仅更新 U”保留两个模块，只替换其中一个规则。论文必须区分模块消融与学习消融。',
 '普通思考对照使用相同公开输入、指导输出接口和每次决策一个额外调用机会；实际 tokens 和耗时可能不同，不能称为实际成本完全相等。',
 '如果仅 U 更新被接受，则 Gf=G₀，某些行会指向同一个规则。按 canonical policy hash 合并运行，并在表注标记别名，不把同一次实验当成多份独立证据。',
 '表 9 的开发轮五组全部采用父规则 G₀、U₀。因此，其“仅 G”和“仅 U”与正式冻结后消融的规则版本可能不同，不能直接并入表 4。'])
final_names=['原生 PAOS','仅 G','仅 U','学习前 G＋U','冻结后 G＋U','仅更新 G 的 RSI','仅更新 U 的 RSI','普通思考＋U']
page('LIBERO 正式物理测试',
 '主指标是完整任务成功率。只有完成规定物理工作、满足顺序与取消约束、保留最终目标并正确结束，才计为成功。每个不同规则组按同一份 48 条件清单运行，不能仅比较各组成功案例。',
 '表 4 LIBERO 正式测试主结果预留表',
 ['方法','成功任务数','成功率及 95% CI','检查点正确率','物理完成进度'],
 [[n,'待测 / 48','待测','待测 / 336','待测 / 960'] for n in final_names],
 [1.7,1.1,1.6,1.3,1.2],[
 '336 = 48 × 7 个计分检查点；960 = 48 × 20 个计划物理段。未到达的检查点与未执行的后续段按任务完成口径计为未完成，同时另报“已到达”和“已尝试”的分母。它们不是 336 或 960 个独立任务样本。',
 'G 的增量贡献：冻结后 G＋U 与仅 U 配对比较；U 的增量贡献：冻结后 G＋U 与仅 G 比较；RSI 收益：冻结后 G＋U 与学习前 G＋U 比较；特定 G 的贡献：冻结后 G＋U 与普通思考＋U 比较。',
 '这四项是已登记的主要比较。报告成功率差（百分点）、配对输赢数量、按 12 种任务历史聚类的 bootstrap 区间和 Holm 校正。随机种子与重采样次数应在分析脚本中冻结；另报逐历史与留一历史分析。48 条件的结果不应当作 48 个完全独立任务家族。',
 '保留决策、协议和物理失败。基础设施中断单独登记；尚有未解决中断时，不把成功数除以“已完成运行数”包装成完整 48 条件结果。'])
page('四类行为诊断',
 '四类诊断来自任务中预先设置的用户事件。它们检验跨会话记忆是否保留了任务状态、指令时效、取消范围和最终目标。下表描述测试方法，不预先断言每一类都已经观察到失败。',
 '表 5 四类诊断的具体操作与计分对象',
 ['问题','具体测试操作','正确行为','应报告指标'],[
 ['中断后恢复','桌面任务先完成放杯子，布丁尚未放置；插入厨房任务并切换会话，随后只说 Continue','恢复尚未完成的放置，保留已完成状态，不重做、不提前结束','正确恢复／已到达恢复机会；另报未到达数与重复动作数'],
 ['采用最新顺序','先给出桌面服务顺序，再连续修改优先级；恢复时不重复完整新顺序','根据最新有效安排选择下一项，同时保留已经完成的工作','最新顺序遵循率；首次顺序违规的位置'],
 ['让取消指令生效','增加一个可选任务，之后明确取消该项，并说明其余工作继续有效','不再执行被取消任务，也不连带取消其他义务','取消后违规执行数；错误扩大取消范围数'],
 ['最终目标与结束','最初约定常规桌面工作完成后仍需执行最终收尾；中间多次打断和归档','先完成前置工作，再完成收尾，最后才结束','收尾时机正确率；遗忘收尾、过早收尾、提前 finish 数']
 ],[1.1,2.2,1.8,1.8],[
 '按原生 PAOS、仅 G、仅 U、完整 G＋U、普通思考＋U 和冻结规则分别汇总。原始轨迹应保留用户指令、归档前后记忆、G 的实际输入输出与 actor 决策，以定位具体失效环节。',
 '未到达的诊断机会不能计作正确，也不能计作该行为已违规。分别报告触达数、正确数、违规数与未到达数，并同时保留完整任务成功率。',
 '同一轨迹可能涉及多种错误，诊断分类不是互斥失败总数。当前检查过的开发记录尚不足以给“取消后仍执行”填写实际失败数，不能为了凑齐四类案例而编造。'])
page('RSI 学习与规则冻结',
 '自动学习从父规则的真实训练轨迹提出局部文本修改，不更新语言模型权重。候选必须先通过结构约束与 critic 审查，再进行同条件回放。保留验证只选择已冻结候选，不提供给提案模型，也不能用来人工修改规则。',
 '表 6 学习阶段四个固定组合的比较表',
 ['回放组合','G 规则','U 规则','训练增量','验证增量','成本约束','接受结果'],[
 ['父规则','G₀','U₀','参照','参照','参照','基线'],
 ['仅 G 编辑','G′','U₀','待完成','待完成','待完成','待定'],
 ['仅 U 编辑','G₀','U′','待完成','待完成','待完成','待定'],
 ['联合编辑','G′','U′','待完成','待完成','待完成','待定']
 ],[1.2,.65,.65,1.1,1.1,1.1,1.0],[
 'G′、U′ 是本轮候选，不等于已接受的 Gf、Uf。本轮两份候选已通过格式校验与 critic；截至取数时学习器仍在运行，尚无最终 report.json。',
 '固定比较规模：4 个训练来源 × 2 次回放 × 4 个不同组合，最多 32 次训练回放；2 个保留验证来源 × 2 次 × 4 组合，最多 16 次验证回放。候选缺失或规则相同时，实际数量按有效组合计算。这些回放不计入表 4 的物理任务数。',
 '接受门槛：训练正确检查点净增大于 0，验证净增不小于 0；完成轨迹数不下降；不存在父规则两次均成功而候选两次均失败的持续退化；相应阶段的逻辑调用不超过父规则的 1.5 倍。满足门槛后按训练增益、编辑算子数、调用数和固定组名顺序选择。',
 '第一次提案尝试没有合格更新：G 达到 8192 输出上限，U 替换文本超过 1200 字符。后续单独登记了输入去重与短局部编辑提示，仍保持硬约束；两版合计最多 4 次算子提案，失败成本计入总学习成本。',
 '回放纠正一个错误不等于真实环境完整成功。只有冻结规则在独立物理测试中优于父规则，才能报告该设置下的 RSI 物理收益；未接受任何更新时保留父规则并如实报告。'])
page('跨环境与基准比较',
 '先在 LIBERO 完成方法开发与规则选择，再冻结方法进入 Casa、Twin 和两个额外 benchmark。各环境使用其明确的任务定义与评分器，不能直接把不同含义的原始分数求平均。',
 '表 7 跨环境主结果预留表',
 ['方法','LIBERO','Casa','Twin','Benchmark 1','Benchmark 2'],
 [[n,'待测','待测','待测','待定名称','待定名称'] for n in final_names],
 [1.7,1.0,1.0,1.0,1.1,1.1],[
 'LIBERO 列使用表 4 的完整任务成功率。Casa、Twin 和两个 benchmark 的任务数、拆分、指标及预算尚待各自正式登记，不能照搬 48 这个分母。已有历史结果不自动属于冻结后的跨环境测试。',
 '“冻结迁移”的含义是沿用已经选定的规则与接口，不根据目标环境测试结果修改规则。如需在新环境训练，应另外划分训练、验证、测试，并单列“目标环境适配”实验。',
 '表中各行是计划覆盖的比较组；相同规则别名合并。两个额外 benchmark 是已要求保留的实验位置，其名称目前未登记，本表不擅自指定。',
 '跨环境结论应逐环境给出配对差值、区间和失败类型；已反复查看过的任务与结果应说明历史接触，不能称为首次未见测试。'])
page('成本与可靠性',
 '对所有正式任务记录实际资源消耗。相同上限控制可用预算，但不会保证所有方法实际花费相同。失败较早的方法可能消耗更少，因此成本必须和任务结果一起报告。',
 '表 8 正式测试的效率与可靠性预留表',
 ['方法','调用／任务','输入／输出 tokens','执行／等待时间','决策／协议／物理失败','未知成本请求'],
 [[n,'待测','待测','待测','待测','待测'] for n in final_names],
 [1.55,1.05,1.2,1.15,1.1,.85],[
 '调用和 tokens：报告所有计划任务的均值与分布，并补充成功任务的条件统计。G、U、actor、优化器与 critic 的调用分别记录；离线学习成本单列，不摊入每个测试任务而隐去。',
 '时间：分别报告总墙钟时间、模型等待、网络等待与物理执行时间，不能把本机断网等待归因于方法推理效率。实际基础设施故障次数和注册技术替代次数另列。',
 '可靠性：协议输出超限或不满足 schema 计为方法协议失败；错误动作或过早结束计为决策失败；物理控制失败单列。若同一任务出现多个症状，终止原因使用互斥主分类，诊断标签可以重叠。',
 'model_cost_complete=false 时，已记录 tokens 只能作为下界，必须同步报告 unknown_requests。不能把未知请求成本当作零，也不能因为缺少完整成本而删除该任务的效果结果。',
 '统一配对条件包括任务历史、初始状态、厨房变体、机器人策略、硬件后端及预算。107 与 108 执行互斥条件分片，同一条件的比较组保留在同一后端；不能把不同硬件的结果直接合并成未分层对照。'])
armnames={'native':'原生 PAOS','g_only':'仅 G₀','u_only':'仅 U₀','gu_parent':'G₀＋U₀','think_u':'普通思考＋U₀'}
status={'success':'成功','decision_failure':'决策失败','protocol_failure':'协议失败'}
rs=[]
for arm in armnames:
 xs=sorted([r for r in data if r['arm']==arm],key=lambda r:r['node'])
 rs.append([armnames[arm]]+[f"{status[x['status']]}\n{x['sessions_completed']}/20 段，{x['correct_checkpoints']}/7 点" for x in xs]+[f"{sum(x['status']=='success' for x in xs)}/2",f"{sum(x['correct_checkpoints'] for x in xs)}/14"])
page('已完成的人工父规则开发结果',
 '以下是真实完成的 v4 开发轮，共 2 个条件、每条件 5 组。两条件都参与方法开发与训练来源构建，因此只能作为开发分析。这里的成功均已核对完整执行、检查点和最终结束。',
 '表 9 v4 开发轮结果 每组两个条件',
 ['方法','条件一结果','条件二结果','完整成功数','正确检查点'],rs,
 [1.45,1.7,1.7,1.05,1.0],[
 '条件一对应 107 的 hEDABC-v1-i0；条件二对应 108 的 hEADCB-v1-i1。ID 仅供审计，正文应使用任务设置的自然语言描述。每个结果中的“段”是已完成物理执行段，“点”是正确计分检查点，二者均不能替代完整任务成功数。',
 '两个仅 G₀ 的失败均为生成输出上限；条件二 G₀＋U₀ 的失败是指导响应出现额外乱码字段，被严格 schema 拒绝。均保留为方法可靠性失败，不按基础设施故障补跑。',
 '这轮开发观察：普通思考＋U₀ 为 2/2，G₀＋U₀ 和仅 U₀ 都为 1/2，原生与仅 G₀ 为 0/2。样本很少，且处于开发阶段，不能证明特定 G 优于普通额外思考，也不能作为 RSI 收益。',
 '此前不同父规则和不同预算版本的结果应单独存档，不与本表合并计算成功率；同样不应择取旧版本中更有利的对照替换本轮失败。'])
rs=[]
for arm in armnames:
 xs=[r for r in data if r['arm']==arm];complete=all(x['model_cost_complete'] for x in xs);unk=sum(x['model_cost_unknown_requests'] for x in xs)
 rs.append([armnames[arm],format(sum(x['model_input_tokens'] for x in xs),','),format(sum(x['model_output_tokens'] for x in xs),','),str(unk),'完整' if complete else '已知下界'])
page('开发成本与表格来源',
 '下表汇总表 9 每组两个开发任务的模型输入与输出 tokens。数值是两任务合计，不是均值，也不是正式测试成本。发生未知成本请求的组只能报告已知下界。',
 '表 10 v4 开发轮模型成本 每组合计两个任务',
 ['方法','输入 tokens','输出 tokens','未知成本请求','成本完整性'],rs,
 [1.7,1.35,1.35,1.25,1.25],[
 '成本完整性来自逐任务 model_cost_complete 与 model_cost_unknown_requests。早停任务执行长度不同，不能仅依据本表的低 token 数认定方法更高效。网络中断后恢复可能留下无法重建的请求成本，这些缺口随原始记录保留。',
 '结果来源：107 和 108 的 study-scope-v4/development-schedule.json，以及十项任务各自的 result.json。设置和拆分来源：study-rsi-v4-sources/registration.json、budget.json。学习约束来源：study-rsi-v4-learning-compact 的控制器登记及 code/learn.py。跨环境范围来源：experiment-allocation.json。',
 '表 1 至表 8 是实验设计与正式结果预留位置；表 9 和表 10 是已完成开发结果。所有“待测”“待完成”“待定”均不是数值结果。正式投稿前应以冻结后导出的结果替换预留项，并保留版本、全部失败、配对条件与成本记录。'])
# Keep the budget erratum with provenance to avoid splitting the setup page.
pages[-1]['notes'].append(pages[1]['notes'].pop())
# Save data model for consistent editable and LaTeX exports.
(OUT/'experiment-tables-data.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2))
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.68);sec.right_margin=Inches(.68)
for sty in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 s=doc.styles[sty];s.font.name='Noto Sans CJK SC';s.font.color.rgb=RGBColor(0,0,0);s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
 s.paragraph_format.space_after=Pt(7)
 s.font.italic=False
 for elem in list(s._element.iter(qn('w:pBdr'))): elem.getparent().remove(elem)
 for elem in s._element.iter(qn('w:color')):
  for a in ['themeColor','themeTint','themeShade']:
   elem.attrib.pop(qn('w:'+a),None)
doc.styles['Normal'].font.size=Pt(10);doc.styles['Normal'].paragraph_format.line_spacing=1.18
doc.styles['Title'].font.size=Pt(20);doc.styles['Heading 1'].font.size=Pt(16);doc.styles['Caption'].font.size=Pt(10);doc.styles['Caption'].font.bold=True
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=foot.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
for idx,p in enumerate(pages):
 if idx:doc.add_page_break()
 doc.add_paragraph('PAOS 论文实验表格' if idx==0 else p['title'],style='Title' if idx==0 else 'Heading 1')
 if idx==0:doc.add_paragraph('2026 年 10 月 24 日  研究设计与结果状态',style='Subtitle')
 doc.add_paragraph(p['intro'])
 doc.add_paragraph(p['caption'],style='Caption')
 t=doc.add_table(rows=1,cols=len(p['headers']));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(p['widths']):t.columns[i].width=Inches(w)
 for i,h in enumerate(p['headers']):t.rows[0].cells[i].text=h
 for row in p['rows']:
  cells=t.add_row().cells
  for i,x in enumerate(row):cells[i].text=str(x)
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();cs=OxmlElement('w:cantSplit');trpr.append(cs)
  if ri==0:
   rep=OxmlElement('w:tblHeader');trpr.append(rep)
  for ci,c in enumerate(row.cells):
   c.width=Inches(p['widths'][ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');bd.append(edge)
   pr.append(bd);mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
   pr.append(mar)
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8ECF0' if ri==0 else 'FFFFFF');pr.append(sh)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(1);para.paragraph_format.line_spacing=1.12
    if idx in [3,5,6,7,8,9] and ci>0:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in para.runs:run.font.size=Pt(9);run.font.bold=ri==0
 doc.add_paragraph()
 for note in p['notes']:doc.add_paragraph(note)
doc.core_properties.title='PAOS 论文实验表格';doc.core_properties.subject='实验设计 正式测试预留表与开发结果';doc.core_properties.author=''
doc.save(OUT/'PAOS-experiment-tables.docx')
# Standalone Chinese XeLaTeX version, with the same data and explanatory notes.
def esc(s):
 s=str(s)
 for a,b in [('₀','0'),('′',"'"),('＋','+'),('×',r'\ensuremath{\times}')]:s=s.replace(a,b)
 for a,b in [('&',r'\&'),('%',r'\%'),('_',r'\_'),('#',r'\#')]:s=s.replace(a,b)
 return s.replace('\n',r'\newline ')
tex=[r'\documentclass[UTF8,10pt,a4paper]{ctexart}',r'\usepackage[margin=1.7cm]{geometry}',r'\usepackage{longtable,array,booktabs}',r'\setlength{\parindent}{0pt}',r'\setlength{\parskip}{6pt}',r'\renewcommand{\arraystretch}{1.25}',r'\begin{document}',r'\title{PAOS 论文实验表格}\date{2026年10月24日}\maketitle']
for i,p in enumerate(pages):
 if i:tex.append(r'\clearpage')
 tex+=[r'\section*{'+esc(p['title'])+'}',esc(p['intro'])]
 # Longtable permits clean wrapping and repeated headers in journal adaptation.
 usable=16.8-(len(p['headers'])-1)*.42
 widths=[usable*w/sum(p['widths']) for w in p['widths']]
 spec=''.join(r'>{\raggedright\arraybackslash}p{'+f'{w:.2f}cm'+'}' for w in widths)
 tex+=[r'\begin{small}',r'\begin{longtable}{'+spec+'}',r'\caption*{'+esc(p['caption'])+r'}\\',r'\toprule',' & '.join(esc(h) for h in p['headers'])+r'\\\midrule',r'\endfirsthead',r'\toprule',' & '.join(esc(h) for h in p['headers'])+r'\\\midrule',r'\endhead']
 tex+=[' & '.join(esc(v) for v in row)+r'\\' for row in p['rows']]
 tex +=[r'\bottomrule',r'\end{longtable}',r'\end{small}']+ [esc(n)+'\n' for n in p['notes']]
tex.append(r'\end{document}')
# caption* requires caption package.
tex.insert(3,r'\usepackage{caption}')
(OUT/'PAOS-experiment-tables.tex').write_text('\n'.join(tex))
print('Created DOCX, LaTeX and source data; 10 tables.')
